from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "manuals" / "Manual_de_implementacion_y_uso_Proctoring_UDS.docx"

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_YELLOW = "FFF4CE"
PALE_RED = "FDE9E7"
GREEN = "2E7D32"
MUTED = "5D6675"
WHITE = "FFFFFF"
BLACK = "172033"


def set_cell_shading(cell, fill):
  tc_pr = cell._tc.get_or_add_tcPr()
  shading = tc_pr.find(qn("w:shd"))
  if shading is None:
    shading = OxmlElement("w:shd")
    tc_pr.append(shading)
  shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
  tc = cell._tc
  tc_pr = tc.get_or_add_tcPr()
  tc_mar = tc_pr.first_child_found_in("w:tcMar")
  if tc_mar is None:
    tc_mar = OxmlElement("w:tcMar")
    tc_pr.append(tc_mar)
  for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
    node = tc_mar.find(qn(f"w:{margin}"))
    if node is None:
      node = OxmlElement(f"w:{margin}")
      tc_mar.append(node)
    node.set(qn("w:w"), str(value))
    node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
  table.alignment = WD_TABLE_ALIGNMENT.LEFT
  table.autofit = False
  table_pr = table._tbl.tblPr
  width = table_pr.find(qn("w:tblW"))
  if width is None:
    width = OxmlElement("w:tblW")
    table_pr.append(width)
  width.set(qn("w:w"), str(sum(widths_dxa)))
  width.set(qn("w:type"), "dxa")
  indent = table_pr.find(qn("w:tblInd"))
  if indent is None:
    indent = OxmlElement("w:tblInd")
    table_pr.append(indent)
  indent.set(qn("w:w"), "120")
  indent.set(qn("w:type"), "dxa")
  grid = table._tbl.tblGrid
  for child in list(grid):
    grid.remove(child)
  for width_value in widths_dxa:
    col = OxmlElement("w:gridCol")
    col.set(qn("w:w"), str(width_value))
    grid.append(col)
  for row in table.rows:
    for index, cell in enumerate(row.cells):
      cell.width = Inches(widths_dxa[index] / 1440)
      tc_pr = cell._tc.get_or_add_tcPr()
      tc_w = tc_pr.find(qn("w:tcW"))
      if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
      tc_w.set(qn("w:w"), str(widths_dxa[index]))
      tc_w.set(qn("w:type"), "dxa")
      cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
      set_cell_margins(cell)


def set_repeat_table_header(row):
  tr_pr = row._tr.get_or_add_trPr()
  header = OxmlElement("w:tblHeader")
  header.set(qn("w:val"), "true")
  tr_pr.append(header)


def set_run_font(run, name="Calibri", size=None, bold=None, italic=None, color=None):
  run.font.name = name
  run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
  run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
  if size is not None:
    run.font.size = Pt(size)
  if bold is not None:
    run.bold = bold
  if italic is not None:
    run.italic = italic
  if color is not None:
    run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(document):
  normal = document.styles["Normal"]
  normal.font.name = "Calibri"
  normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
  normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
  normal.font.size = Pt(11)
  normal.font.color.rgb = RGBColor.from_string(BLACK)
  normal.paragraph_format.space_after = Pt(6)
  normal.paragraph_format.line_spacing = 1.25

  for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
  ):
    style = document.styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

  for list_name in ("List Bullet", "List Number"):
    style = document.styles[list_name]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.left_indent = Inches(0.375)
    style.paragraph_format.first_line_indent = Inches(-0.188)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25

  code = document.styles.add_style("Code Block", 1)
  code.font.name = "Consolas"
  code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
  code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
  code.font.size = Pt(9)
  code.font.color.rgb = RGBColor.from_string(BLACK)
  code.paragraph_format.space_before = Pt(4)
  code.paragraph_format.space_after = Pt(8)
  code.paragraph_format.line_spacing = 1.0
  code.paragraph_format.keep_together = True


def add_page_number(paragraph):
  paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
  run = paragraph.add_run("Página ")
  set_run_font(run, size=9, color=MUTED)
  begin = OxmlElement("w:fldChar")
  begin.set(qn("w:fldCharType"), "begin")
  instr = OxmlElement("w:instrText")
  instr.set(qn("xml:space"), "preserve")
  instr.text = " PAGE "
  separate = OxmlElement("w:fldChar")
  separate.set(qn("w:fldCharType"), "separate")
  end = OxmlElement("w:fldChar")
  end.set(qn("w:fldCharType"), "end")
  run._r.extend([begin, instr, separate, end])


def configure_sections(document):
  for section in document.sections:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("PROCTORING UDS  |  MANUAL DE IMPLEMENTACIÓN Y USO")
    set_run_font(run, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_p = footer.paragraphs[0]
    add_page_number(footer_p)


def add_title_page(document):
  for _ in range(6):
    document.add_paragraph()
  p = document.add_paragraph()
  p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.space_after = Pt(10)
  run = p.add_run("MANUAL DE IMPLEMENTACIÓN Y USO")
  set_run_font(run, size=26, bold=True, color=NAVY)

  p = document.add_paragraph()
  p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.space_after = Pt(28)
  run = p.add_run("Sistema de Proctoring UDS integrado con Moodle")
  set_run_font(run, size=16, color=BLUE)

  p = document.add_paragraph()
  p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  run = p.add_run("Instalación de MariaDB, API independiente, panel de revisión, evidencia y operación académica")
  set_run_font(run, size=11, italic=True, color=MUTED)

  for _ in range(8):
    document.add_paragraph()
  p = document.add_paragraph()
  p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.space_after = Pt(3)
  set_run_font(p.add_run("Versión del manual: 1.1"), size=10, bold=True, color=NAVY)
  p = document.add_paragraph()
  p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  set_run_font(p.add_run("Entorno objetivo: servidor de pruebas y piloto institucional"), size=10, color=MUTED)
  document.add_page_break()


def add_paragraph(document, text, bold_prefix=None):
  p = document.add_paragraph()
  if bold_prefix and text.startswith(bold_prefix):
    set_run_font(p.add_run(bold_prefix), bold=True)
    set_run_font(p.add_run(text[len(bold_prefix):]))
  else:
    set_run_font(p.add_run(text))
  return p


def add_bullets(document, items):
  for item in items:
    p = document.add_paragraph(style="List Bullet")
    set_run_font(p.add_run(item))


def add_steps(document, items):
  for index, item in enumerate(items, start=1):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    set_run_font(p.add_run(f"{index}.  "))
    set_run_font(p.add_run(item))


def add_code(document, code):
  table = document.add_table(rows=1, cols=1)
  set_table_geometry(table, [9360])
  cell = table.cell(0, 0)
  set_cell_shading(cell, LIGHT_GRAY)
  p = cell.paragraphs[0]
  p.style = document.styles["Code Block"]
  for index, line in enumerate(code.strip().splitlines()):
    if index:
      p.add_run().add_break()
    run = p.add_run(line)
    set_run_font(run, name="Consolas", size=9, color=BLACK)
  document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_callout(document, label, text, kind="info"):
  fill = {"info": LIGHT_BLUE, "warning": PALE_YELLOW, "danger": PALE_RED}.get(kind, LIGHT_BLUE)
  table = document.add_table(rows=1, cols=1)
  set_table_geometry(table, [9360])
  cell = table.cell(0, 0)
  set_cell_shading(cell, fill)
  p = cell.paragraphs[0]
  p.paragraph_format.space_after = Pt(0)
  set_run_font(p.add_run(f"{label}: "), bold=True, color=NAVY)
  set_run_font(p.add_run(text), color=BLACK)
  document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(document, headers, rows, widths):
  table = document.add_table(rows=1, cols=len(headers))
  set_table_geometry(table, widths)
  set_repeat_table_header(table.rows[0])
  for index, header in enumerate(headers):
    cell = table.rows[0].cells[index]
    set_cell_shading(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(header) < 18 else WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run(header), bold=True, color=NAVY, size=9.5)
  for row_data in rows:
    cells = table.add_row().cells
    for index, value in enumerate(row_data):
      p = cells[index].paragraphs[0]
      p.paragraph_format.space_after = Pt(0)
      p.alignment = WD_ALIGN_PARAGRAPH.LEFT
      set_run_font(p.add_run(str(value)), size=9.5)
  set_table_geometry(table, widths)
  document.add_paragraph().paragraph_format.space_after = Pt(2)
  return table


def add_toc(document):
  document.add_heading("Contenido", level=1)
  entries = [
    "1. Introducción y alcance",
    "2. Arquitectura y componentes",
    "3. Requisitos previos",
    "4. Instalación de MariaDB",
    "5. Preparación del almacenamiento de evidencia",
    "6. Generación y administración de secretos",
    "7. Instalación y configuración de la API",
    "8. Instalación de la aplicación web y panel",
    "9. Publicación HTTPS y dominios",
    "10. Instalación de los plugins en Moodle",
    "11. Configuración y permisos en Moodle",
    "12. Configuración de cuestionarios protegidos",
    "13. Uso por el estudiante",
    "14. Uso por docentes y revisores",
    "15. Evidencia, auditoría y retención",
    "16. Mantenimiento, copias y actualizaciones",
    "17. Pruebas y aceptación",
    "18. Solución de problemas",
    "19. Anexos de referencia",
  ]
  table = document.add_table(rows=10, cols=2)
  set_table_geometry(table, [4680, 4680])
  for index, entry in enumerate(entries):
    cell = table.cell(index % 10, index // 10)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(entry), color=DARK_BLUE, size=10)


def build_document():
  document = Document()
  configure_styles(document)
  configure_sections(document)
  add_title_page(document)

  document.add_heading("Control del documento", level=1)
  add_table(document, ["Campo", "Valor"], [
    ("Documento", "Manual de implementación y uso - Proctoring UDS"),
    ("Versión", "1.1"),
    ("Audiencia", "Administradores de servidor, administradores Moodle, docentes, revisores y soporte"),
    ("Alcance", "Servidor de pruebas, validación funcional y preparación del piloto"),
    ("Base externa", "MariaDB exclusiva para proctoring"),
  ], [2700, 6660])
  add_callout(document, "Principio de separación", "Moodle conserva usuarios, cursos, cuestionarios, intentos y calificaciones. La plataforma de proctoring conserva sus propias sesiones, eventos, alertas y evidencia. La API no consulta tablas de Moodle.")
  add_toc(document)

  document.add_heading("1. Introducción y alcance", level=1)
  add_paragraph(document, "Este manual explica cómo implementar y operar el sistema de Proctoring UDS integrado con Moodle. Incluye la instalación de la base MariaDB independiente, la API, la aplicación web, el panel de revisión, los plugins Moodle y el almacenamiento de evidencia.")
  add_paragraph(document, "El sistema acompaña exámenes protegidos mediante verificación previa de cámara, detección local de rostro, prueba de vida, monitoreo de incidencias y revisión humana posterior. Las alertas nunca modifican automáticamente una calificación.")
  document.add_heading("1.1 Usuarios del manual", level=2)
  add_table(document, ["Rol", "Responsabilidad"], [
    ("Administrador de servidor", "Instala MariaDB, API, web, certificados, servicios y almacenamiento."),
    ("Administrador Moodle", "Instala plugins, configura URLs, secretos, capacidades y cuestionarios."),
    ("Docente", "Configura el cuestionario y consulta sesiones de sus cursos."),
    ("Revisor", "Analiza alertas, cronología y evidencia autorizada; registra una decisión humana."),
    ("Estudiante", "Completa la preparación y mantiene la cámara disponible durante el intento."),
    ("Soporte", "Diagnostica conectividad, cámara, estados y correlación de registros."),
  ], [2500, 6860])
  document.add_heading("1.2 Límites funcionales", level=2)
  add_bullets(document, [
    "No graba vídeo continuo.",
    "No infiere emociones.",
    "No modifica intentos ni calificaciones Moodle.",
    "No copia nombres ni correos a la plataforma externa; utiliza identificadores técnicos.",
    "La evidencia biométrica requiere una capacidad explícita y cada acceso queda auditado.",
  ])

  document.add_heading("2. Arquitectura y componentes", level=1)
  add_code(document, """
Navegador del estudiante
  |-- Moodle
  |-- Aplicación web de preparación y monitor

Moodle -- HTTPS + clave de integración --> API Proctoring
                                              |-- MariaDB propia
                                              |-- almacenamiento S3 privado
                                              `-- registros y auditoría

Moodle -- token SSO breve --> API --> cookie segura --> Panel de revisión
""")
  document.add_heading("2.1 Componentes", level=2)
  add_table(document, ["Componente", "Función", "Puerto interno sugerido"], [
    ("Moodle", "Gestiona identidad, cursos, intentos y notas.", "80/443 mediante servidor web"),
    ("API", "Sesiones, tokens, eventos, alertas, evidencia y autorización.", "3001"),
    ("Aplicación web", "Preparación del estudiante, monitor y panel.", "3000"),
    ("MariaDB", "Persistencia exclusiva de proctoring.", "3306 privado"),
    ("S3 compatible", "Objetos de evidencia cifrados.", "HTTPS privado"),
  ], [2300, 4760, 2300])
  document.add_heading("2.2 Dominios recomendados", level=2)
  add_bullets(document, [
    "Moodle: https://moodle-pruebas.institucion.edu",
    "Aplicación y panel: https://proctoring-pruebas.institucion.edu",
    "API: https://api-proctoring-pruebas.institucion.edu",
  ])
  add_callout(document, "Importante", "La cámara, las cookies seguras y la comunicación entre servicios requieren HTTPS con certificados válidos.", "warning")

  document.add_heading("3. Requisitos previos", level=1)
  document.add_heading("3.1 Software", level=2)
  add_bullets(document, [
    "Moodle 4.3.3 o superior y PHP compatible con la versión instalada.",
    "Node.js 22 o superior y pnpm 11 o superior.",
    "MariaDB 10.11.14 o superior.",
    "Nginx o Apache para publicación HTTPS.",
    "Git para descargar y actualizar el código.",
    "Almacenamiento de objetos compatible con S3 y bucket privado.",
  ])
  document.add_heading("3.2 Recursos para pruebas", level=2)
  add_table(document, ["Recurso", "Mínimo recomendado"], [
    ("CPU", "4 núcleos"),
    ("Memoria", "8 GB"),
    ("Disco de aplicaciones y logs", "40 GB"),
    ("Base MariaDB", "20 GB iniciales con monitoreo de crecimiento"),
    ("Red", "HTTPS estable entre Moodle, API, panel y clientes"),
  ], [3400, 5960])
  document.add_heading("3.3 Datos que debe definir la institución", level=2)
  add_bullets(document, [
    "Plazo de retención de evidencia.",
    "Roles autorizados para evidencia biométrica.",
    "Procedimiento de revisión y tratamiento de falsos positivos.",
    "Canal de soporte e incidentes.",
    "Dominios públicos y certificados.",
    "Responsables de MariaDB, almacenamiento y Moodle.",
  ])

  document.add_heading("4. Instalación de MariaDB", level=1)
  document.add_heading("4.1 Instalar el servicio", level=2)
  add_paragraph(document, "Instale MariaDB mediante el gestor de paquetes de su sistema operativo. El servicio puede ejecutarse en el mismo servidor de pruebas que Moodle, siempre que utilice una base, un usuario y permisos separados.")
  add_code(document, """
sudo apt update
sudo apt install mariadb-server mariadb-client
sudo systemctl enable --now mariadb
sudo systemctl status mariadb
""")
  document.add_heading("4.2 Crear usuario y base", level=2)
  add_code(document, """
sudo mariadb

CREATE DATABASE proctoring CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;
CREATE USER 'proctoring_app'@'127.0.0.1' IDENTIFIED BY 'CONTRASENA_LARGA_Y_ALEATORIA';
GRANT SELECT, INSERT, UPDATE, DELETE, CREATE, ALTER, INDEX, REFERENCES, DROP
  ON proctoring.* TO 'proctoring_app'@'127.0.0.1';
FLUSH PRIVILEGES;
exit
""")
  add_callout(document, "Separación obligatoria", "No utilice el usuario, la contraseña ni la base de datos de Moodle. No conceda al usuario proctoring_app acceso a la base Moodle.", "danger")
  document.add_heading("4.3 Verificar conexión", level=2)
  add_code(document, """
mariadb --host=127.0.0.1 --user=proctoring_app --password proctoring -e "SELECT 1;"
""")
  document.add_heading("4.4 Acceso de red", level=2)
  add_paragraph(document, "Si la API y MariaDB están en el mismo servidor, escuche solo en 127.0.0.1. Si están separados, permita exclusivamente la IP privada de la API, use TLS y bloquee el puerto 3306 desde Internet.")

  document.add_heading("5. Preparación del almacenamiento de evidencia", level=1)
  add_steps(document, [
    "Cree un bucket exclusivo, por ejemplo proctoring-evidence-test.",
    "Bloquee el acceso público del bucket.",
    "Cree una identidad de servicio exclusiva para la API.",
    "Conceda únicamente lectura, escritura y eliminación dentro del bucket seleccionado.",
    "Habilite HTTPS y cifrado del lado del almacenamiento.",
    "Registre endpoint, región, bucket, clave de acceso y secreto en un gestor seguro.",
  ])
  add_table(document, ["Variable", "Descripción"], [
    ("S3_ENDPOINT", "URL HTTPS del servicio compatible con S3."),
    ("S3_REGION", "Región configurada para el bucket."),
    ("S3_BUCKET", "Nombre del bucket privado."),
    ("S3_ACCESS_KEY_ID", "Identificador de la credencial de la API."),
    ("S3_SECRET_ACCESS_KEY", "Secreto de la credencial de la API."),
    ("S3_FORCE_PATH_STYLE", "true cuando el proveedor requiere rutas por prefijo de bucket."),
  ], [3100, 6260])
  add_callout(document, "Protección de evidencia", "La API cifra cada captura con AES-256-GCM antes de almacenarla. El bucket nunca debe ser público y las URLs emitidas por la API caducan rápidamente.")

  document.add_heading("6. Generación y administración de secretos", level=1)
  add_code(document, """
openssl rand -base64 48   # MOODLE_INTEGRATION_KEY
openssl rand -base64 48   # JWT_SECRET
openssl rand -base64 48   # PANEL_SSO_SECRET
openssl rand -base64 32   # EVIDENCE_ENCRYPTION_KEY
""")
  add_table(document, ["Secreto", "Ubicación", "Finalidad"], [
    ("MOODLE_INTEGRATION_KEY", "Moodle y API", "Autentica llamadas servidor a servidor."),
    ("JWT_SECRET", "Solo API", "Firma tokens temporales de navegador y panel."),
    ("PANEL_SSO_SECRET", "Moodle y API", "Firma el acceso SSO del revisor."),
    ("EVIDENCE_ENCRYPTION_KEY", "Solo API", "Cifra y descifra evidencia retenida."),
    ("Credenciales S3", "Solo API", "Acceso restringido al bucket."),
    ("DATABASE_URL", "Solo API", "Acceso a la base externa."),
  ], [2600, 1900, 4860])
  add_callout(document, "Regla", "No incluya secretos en Git, correos, capturas, tickets o documentación. Proteja los archivos de entorno con permisos 600.", "danger")

  document.add_heading("7. Instalación y configuración de la API", level=1)
  document.add_heading("7.1 Descargar el proyecto", level=2)
  add_code(document, """
sudo mkdir -p /opt/proctoring-uds
sudo chown "$USER":"$USER" /opt/proctoring-uds
git clone URL_DEL_REPOSITORIO /opt/proctoring-uds
cd /opt/proctoring-uds
pnpm install --frozen-lockfile
""")
  document.add_heading("7.2 Crear el archivo de entorno", level=2)
  add_code(document, """
sudo mkdir -p /etc/proctoring
sudo nano /etc/proctoring/api.env
sudo chmod 600 /etc/proctoring/api.env
""")
  add_code(document, """
PROCTORING_HOST=127.0.0.1
PROCTORING_PORT=3001
DATABASE_URL=mysql://proctoring_app:CAMBIAR@127.0.0.1:3306/proctoring
MOODLE_INTEGRATION_KEY=CAMBIAR
JWT_SECRET=CAMBIAR
PANEL_SSO_SECRET=CAMBIAR
EVIDENCE_ENCRYPTION_KEY=CAMBIAR
WEB_ORIGIN=https://proctoring-pruebas.institucion.edu
API_PUBLIC_URL=https://api-proctoring-pruebas.institucion.edu
PANEL_URL=https://proctoring-pruebas.institucion.edu
EVIDENCE_RETENTION_DAYS=30
S3_ENDPOINT=https://s3.institucion.edu
S3_REGION=us-east-1
S3_BUCKET=proctoring-evidence-test
S3_ACCESS_KEY_ID=CAMBIAR
S3_SECRET_ACCESS_KEY=CAMBIAR
S3_FORCE_PATH_STYLE=true
""")
  document.add_heading("7.3 Ejecutar migraciones", level=2)
  add_code(document, """
cd /opt/proctoring-uds
set -a
. /etc/proctoring/api.env
set +a
pnpm --filter @proctoring/api migrate
""")
  add_paragraph(document, "Confirme que se aplicaron las migraciones 001, 002, 003 y 004. No edite manualmente las tablas creadas.")
  document.add_heading("7.4 Crear el servicio", level=2)
  add_code(document, """
[Unit]
Description=Proctoring UDS API
After=network.target mariadb.service

[Service]
Type=simple
User=www-data
WorkingDirectory=/opt/proctoring-uds
EnvironmentFile=/etc/proctoring/api.env
ExecStart=/usr/bin/pnpm api:start
Restart=on-failure
RestartSec=5
NoNewPrivileges=true
PrivateTmp=true

[Install]
WantedBy=multi-user.target
""")
  add_code(document, """
sudo systemctl daemon-reload
sudo systemctl enable --now proctoring-api
sudo systemctl status proctoring-api
curl http://127.0.0.1:3001/health
""")
  add_paragraph(document, "La respuesta normal es {\"status\":\"ok\",\"database\":\"available\"}. Un código 503 indica que la API no puede consultar MariaDB.")

  document.add_heading("8. Instalación de la aplicación web y panel", level=1)
  document.add_heading("8.1 Configuración", level=2)
  add_code(document, """
NEXT_PUBLIC_PROCTORING_API_URL=https://api-proctoring-pruebas.institucion.edu
NEXT_PUBLIC_EVIDENCE_INTERVAL_SECONDS=60
MOODLE_ORIGIN=https://moodle-pruebas.institucion.edu
""")
  document.add_heading("8.2 Compilar", level=2)
  add_code(document, """
cd /opt/proctoring-uds
set -a
. /etc/proctoring/web.env
set +a
pnpm web:build
""")
  document.add_heading("8.3 Servicio web", level=2)
  add_code(document, """
[Unit]
Description=Proctoring UDS Web
After=network.target proctoring-api.service

[Service]
Type=simple
User=www-data
WorkingDirectory=/opt/proctoring-uds
EnvironmentFile=/etc/proctoring/web.env
ExecStart=/usr/bin/pnpm --filter @proctoring/web start
Restart=on-failure
RestartSec=5
NoNewPrivileges=true
PrivateTmp=true

[Install]
WantedBy=multi-user.target
""")
  add_code(document, """
sudo systemctl daemon-reload
sudo systemctl enable --now proctoring-web
sudo systemctl status proctoring-web
""")

  document.add_heading("9. Publicación HTTPS y dominios", level=1)
  document.add_heading("9.1 Proxy de la API", level=2)
  add_code(document, """
server {
  listen 443 ssl http2;
  server_name api-proctoring-pruebas.institucion.edu;
  client_max_body_size 512k;

  location / {
    proxy_pass http://127.0.0.1:3001;
    proxy_http_version 1.1;
    proxy_set_header Host $host;
    proxy_set_header X-Real-IP $remote_addr;
    proxy_set_header X-Forwarded-For $proxy_add_x_forwarded_for;
    proxy_set_header X-Forwarded-Proto https;
  }
}
""")
  document.add_heading("9.2 Proxy de la aplicación web", level=2)
  add_code(document, """
server {
  listen 443 ssl http2;
  server_name proctoring-pruebas.institucion.edu;

  location / {
    proxy_pass http://127.0.0.1:3000;
    proxy_http_version 1.1;
    proxy_set_header Host $host;
    proxy_set_header X-Real-IP $remote_addr;
    proxy_set_header X-Forwarded-For $proxy_add_x_forwarded_for;
    proxy_set_header X-Forwarded-Proto https;
  }
}
""")
  add_code(document, """
sudo nginx -t
sudo systemctl reload nginx
curl -i https://api-proctoring-pruebas.institucion.edu/health
""")
  add_callout(document, "Origen exacto", "WEB_ORIGIN debe coincidir exactamente con el origen público del panel. MOODLE_ORIGIN debe coincidir exactamente con Moodle. Cambiar una URL requiere recompilar o reiniciar el componente correspondiente.", "warning")

  document.add_heading("10. Instalación de los plugins en Moodle", level=1)
  document.add_heading("10.1 Copiar los plugins", level=2)
  add_code(document, """
sudo cp -R apps/moodle/local/proctoring /RUTA_MOODLE/local/proctoring
sudo cp -R apps/moodle/mod/quiz/accessrule/proctoring \
  /RUTA_MOODLE/mod/quiz/accessrule/proctoring

sudo chown -R www-data:www-data /RUTA_MOODLE/local/proctoring
sudo chown -R www-data:www-data \
  /RUTA_MOODLE/mod/quiz/accessrule/proctoring
""")
  document.add_heading("10.2 Actualizar Moodle", level=2)
  add_code(document, """
sudo -u www-data php /RUTA_MOODLE/admin/cli/upgrade.php --non-interactive
sudo -u www-data php /RUTA_MOODLE/admin/cli/purge_caches.php
""")
  add_paragraph(document, "Instale primero local/proctoring y después quizaccess_proctoring. Moodle conservará únicamente la referencia remota, el estado y la caducidad de cada sesión.")

  document.add_heading("11. Configuración y permisos en Moodle", level=1)
  document.add_heading("11.1 Configuración general", level=2)
  add_steps(document, [
    "Ingrese como administrador Moodle.",
    "Abra Administración del sitio > Plugins > Plugins locales > Proctoring.",
    "Configure la URL HTTPS pública de la API.",
    "Pegue MOODLE_INTEGRATION_KEY, idéntica a la API.",
    "Configure la URL HTTPS del panel.",
    "Pegue PANEL_SSO_SECRET, idéntico a la API.",
    "Guarde y purgue cachés.",
  ])
  document.add_heading("11.2 Capacidades", level=2)
  add_table(document, ["Capacidad", "Uso recomendado"], [
    ("local/proctoring:viewowncoursereports", "Docentes: consultar sus cursos."),
    ("local/proctoring:reviewowncoursealerts", "Docentes o revisores: decidir alertas de sus cursos."),
    ("local/proctoring:viewinstitutionreports", "Coordinadores: alcance institucional."),
    ("local/proctoring:viewbiometricevidence", "Rol restringido: visualizar evidencia."),
    ("local/proctoring:managepolicies", "Administradores autorizados: políticas."),
  ], [4700, 4660])
  add_callout(document, "Mínimo privilegio", "No conceda acceso biométrico automáticamente a docentes ni administradores técnicos. Cree un rol específico y documente cada asignación.", "warning")

  document.add_heading("12. Configuración de cuestionarios protegidos", level=1)
  add_steps(document, [
    "Abra el curso y active el modo edición.",
    "Cree o edite un cuestionario.",
    "Abra las restricciones de acceso del cuestionario.",
    "Active Require proctoring.",
    "Seleccione browser, seb o either.",
    "Seleccione la política ante indisponibilidad.",
    "Guarde y realice un intento con una cuenta estudiante de prueba.",
  ])
  add_table(document, ["Modalidad", "Comportamiento"], [
    ("browser", "Flujo para escritorio y móvil; no exige Safe Exam Browser."),
    ("seb", "Requiere que Moodle tenga activa su regla nativa Safe Exam Browser."),
    ("either", "Usa seb cuando Moodle lo exige; en caso contrario usa browser."),
  ], [2200, 7160])
  add_callout(document, "Safe Exam Browser", "La integración de proctoring no reemplaza la validación nativa de Moodle. Configure y pruebe ambas reglas en el cuestionario.")

  document.add_heading("13. Uso por el estudiante", level=1)
  document.add_heading("13.1 Antes del examen", level=2)
  add_bullets(document, [
    "Usar un navegador actualizado y una cámara funcional.",
    "Cerrar aplicaciones que utilicen la cámara.",
    "Disponer de iluminación frontal y conexión estable.",
    "Mantener un solo rostro visible y retirar objetos que bloqueen la cámara.",
    "Conceder permiso de cámara cuando el navegador lo solicite.",
  ])
  document.add_heading("13.2 Preparación", level=2)
  add_steps(document, [
    "Iniciar el intento desde Moodle.",
    "Esperar la apertura de la preparación externa.",
    "Aceptar el permiso de cámara.",
    "Ubicar el rostro dentro del encuadre y permanecer solo frente a la cámara.",
    "Completar dos acciones aleatorias: parpadeo o giro a izquierda/derecha.",
    "Esperar la confirmación y el retorno a Moodle.",
    "Realizar el examen sin cerrar el monitor incorporado.",
  ])
  document.add_heading("13.3 Durante el intento", level=2)
  add_paragraph(document, "El sistema puede registrar interrupción de cámara, ausencia o múltiples rostros, encuadre incorrecto, cambio de visibilidad y desconexión. Estos eventos producen información revisable; no representan por sí solos una sanción.")
  document.add_heading("13.4 Problemas frecuentes del estudiante", level=2)
  add_table(document, ["Problema", "Acción"], [
    ("Cámara denegada", "Permitir cámara en el sitio, cerrar otras aplicaciones y recargar."),
    ("No se detecta rostro", "Mejorar iluminación, mirar al frente y limpiar la lente."),
    ("Dos rostros", "Permanecer solo dentro del campo de la cámara."),
    ("Conexión interrumpida", "No cerrar Moodle; esperar reconexión. La cola reintentará los eventos."),
    ("Token vencido", "Volver al cuestionario e iniciar nuevamente el acceso."),
  ], [2900, 6460])

  document.add_heading("14. Uso por docentes y revisores", level=1)
  document.add_heading("14.1 Abrir el panel", level=2)
  add_steps(document, [
    "Ingresar a Moodle con la cuenta institucional.",
    "Seleccionar Proctoring review en la navegación.",
    "Moodle emitirá un acceso breve con cursos y capacidades vigentes.",
    "La API creará una sesión segura del panel y mostrará únicamente el alcance autorizado.",
  ])
  document.add_heading("14.2 Revisar una sesión", level=2)
  add_steps(document, [
    "Seleccionar una sesión por curso e intento.",
    "Revisar modalidad, estado y número de alertas abiertas.",
    "Leer la cronología completa antes de emitir una decisión.",
    "Consultar evidencia solo si posee capacidad biométrica y existe una necesidad legítima.",
    "Marcar la alerta como revisada o descartada.",
    "Agregar una nota objetiva, sin diagnósticos ni inferencias emocionales.",
  ])
  add_callout(document, "Decisión humana", "Una alerta indica una condición técnica o visual que requiere contexto. Nunca cambie una nota únicamente por la existencia de una alerta.", "danger")
  document.add_heading("14.3 Aislamiento por curso", level=2)
  add_paragraph(document, "El servidor filtra las sesiones por los cursos incluidos en el token Moodle. Modificar una URL o UUID no debe permitir acceder a otro curso. Un administrador institucional requiere una capacidad global explícita.")

  document.add_heading("15. Evidencia, auditoría y retención", level=1)
  document.add_heading("15.1 Datos conservados", level=2)
  add_bullets(document, [
    "Captura de referencia de identidad.",
    "Capturas por intervalo configurable.",
    "Capturas asociadas a alertas, con límite de frecuencia.",
    "Eventos técnicos y estados de revisión.",
    "Auditoría de visualización, descarga y eliminación.",
  ])
  document.add_heading("15.2 Acceso", level=2)
  add_paragraph(document, "La API comprueba sesión del panel, capacidad biométrica y curso autorizado. Después emite una URL temporal de 60 segundos. La lectura descifra la evidencia en memoria y registra la descarga.")
  document.add_heading("15.3 Eliminación automática", level=2)
  add_code(document, """
cd /opt/proctoring-uds
set -a
. /etc/proctoring/api.env
set +a
pnpm api:evidence:purge
""")
  add_paragraph(document, "Programe la tarea al menos una vez al día. La tarea elimina primero el objeto, registra retention_delete y marca el metadato como eliminado.")

  document.add_heading("16. Mantenimiento, copias y actualizaciones", level=1)
  document.add_heading("16.1 Monitoreo diario", level=2)
  add_bullets(document, [
    "Estado de proctoring-api y proctoring-web.",
    "Respuesta de /health.",
    "Espacio y conexiones de MariaDB.",
    "Errores y capacidad del bucket S3.",
    "Ejecución de la tarea de retención.",
    "Certificados próximos a vencer.",
  ])
  document.add_heading("16.2 Copia de MariaDB", level=2)
  add_code(document, """
mariadb-dump --host=127.0.0.1 --user=proctoring_app --password \
  proctoring \
  > proctoring_$(date +%Y%m%d).dump
""")
  add_paragraph(document, "Proteja y cifre las copias. La copia de MariaDB no contiene los objetos del bucket; el plan de recuperación debe considerar ambos componentes y conservar vencimientos.")
  document.add_heading("16.3 Actualización", level=2)
  add_steps(document, [
    "Realizar copia de MariaDB y registrar la versión actual.",
    "Detener nuevas evaluaciones durante la ventana.",
    "Descargar la versión aprobada del repositorio.",
    "Ejecutar pnpm install --frozen-lockfile y pnpm --filter @proctoring/api migrate.",
    "Compilar la web con pnpm web:build.",
    "Actualizar plugins Moodle y ejecutar admin/cli/upgrade.php.",
    "Reiniciar servicios, verificar /health y ejecutar una sesión sintética.",
  ])
  document.add_heading("16.4 Rotación de secretos", level=2)
  add_paragraph(document, "Rote secretos durante una ventana sin exámenes. La clave de cifrado requiere conservar versiones anteriores o recifrar la evidencia retenida; no la sustituya mientras existan objetos que dependan de ella sin un procedimiento aprobado.")

  document.add_heading("17. Pruebas y aceptación", level=1)
  document.add_heading("17.1 Pruebas automáticas", level=2)
  add_code(document, """
cd /opt/proctoring-uds
pnpm test
pnpm web:build
""")
  add_code(document, """
cd /RUTA_MOODLE
sudo -u www-data vendor/bin/phpunit --testsuite local_proctoring
sudo -u www-data vendor/bin/phpunit --testsuite quizaccess_proctoring
""")
  document.add_heading("17.2 Prueba funcional", level=2)
  add_steps(document, [
    "Crear dos cursos, dos docentes limitados y estudiantes de prueba.",
    "Completar preparación en escritorio, Android e iOS.",
    "Completar un flujo con la regla nativa Safe Exam Browser.",
    "Generar alertas de cámara, rostro, pestaña y red.",
    "Confirmar que el docente A nunca recibe información del curso B.",
    "Confirmar acceso denegado a evidencia sin capacidad biométrica.",
    "Revisar una alerta y verificar que la calificación Moodle no cambió.",
  ])
  document.add_heading("17.3 Prueba de 1.000 sesiones", level=2)
  add_code(document, """
export PILOT_API_URL=https://api-proctoring-pruebas.institucion.edu
export MOODLE_INTEGRATION_KEY='SECRETO'
export PILOT_SESSIONS=1000
export PILOT_CONCURRENCY=50
export PILOT_MAX_RETRIES=2
export PILOT_SEED=prueba-001
pnpm pilot:load
""")
  add_table(document, ["Métrica", "Umbral inicial"], [
    ("Fallos terminales", "0"),
    ("Latencia p95", "Menor de 1.000 ms"),
    ("Operaciones reintentadas", "Menos de 1%"),
    ("Acceso cruzado", "0 casos"),
    ("Errores MariaDB/S3", "0 durante la ventana"),
  ], [3900, 5460])

  document.add_heading("18. Solución de problemas", level=1)
  add_table(document, ["Síntoma", "Comprobación", "Solución"], [
    ("/health devuelve 503", "Estado y DATABASE_URL de MariaDB.", "Restablecer conexión; no mostrar credenciales en tickets."),
    ("Moodle no crea sesión", "URL API, clave compartida, TLS y conectividad saliente.", "Corregir configuración y probar curl desde el servidor Moodle."),
    ("El panel muestra acceso vencido", "Hora del servidor y secreto SSO.", "Sincronizar reloj y volver a abrir desde Moodle."),
    ("El docente no ve sesiones", "Capacidad y matrícula del curso.", "Asignar rol correcto y emitir un acceso nuevo."),
    ("Evidencia no disponible", "Capacidad biométrica, curso y estado S3.", "Corregir autorización o almacenamiento; no hacer público el bucket."),
    ("La cámara no abre", "HTTPS, permisos del navegador y otra aplicación usando cámara.", "Conceder permiso, cerrar aplicaciones y recargar."),
    ("Eventos pendientes", "Conectividad del cliente y disponibilidad API.", "Mantener la página abierta y restablecer la red."),
    ("Cierre remoto falla", "Logs API y conectividad al finalizar.", "Investigar close_failed; no impedir la entrega ni cambiar notas."),
  ], [2200, 3100, 4060])
  document.add_heading("18.1 Registros útiles", level=2)
  add_code(document, """
sudo journalctl -u proctoring-api -n 200 --no-pager
sudo journalctl -u proctoring-web -n 200 --no-pager
sudo systemctl status mariadb
curl -i https://api-proctoring-pruebas.institucion.edu/health
""")
  add_callout(document, "Soporte seguro", "Solicite intento, hora, dispositivo y X-Correlation-ID. No solicite evidencia, claves o contraseñas por correo o chat.", "warning")

  document.add_heading("19. Anexos de referencia", level=1)
  document.add_heading("19.1 Lista previa a la puesta en marcha", level=2)
  checklist = [
    "MariaDB separada y usuario de mínimo privilegio.",
    "Migraciones 001-004 aplicadas.",
    "Bucket S3 privado y credencial exclusiva.",
    "Dominios y certificados válidos.",
    "API y web administradas como servicios.",
    "Plugins Moodle instalados y actualizados.",
    "URLs y secretos coincidentes entre Moodle y API.",
    "Roles por curso y capacidad biométrica revisados.",
    "Retención diaria programada.",
    "Copias y restauración probadas.",
    "Pruebas automáticas y PHPUnit aprobadas.",
    "Escritorio, móvil y Safe Exam Browser validados.",
    "Carga de 1.000 sesiones dentro de umbrales.",
    "Canales de soporte, privacidad e incidentes definidos.",
  ]
  for item in checklist:
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    set_run_font(p.add_run("☐ "), size=12, color=BLUE)
    set_run_font(p.add_run(item))

  document.add_heading("19.2 Prueba híbrida con Moodle remoto", level=2)
  add_paragraph(document, "Para una validación temporal, los plugins pueden instalarse en un Moodle remoto mientras API, web, MariaDB y almacenamiento se ejecutan en una computadora de desarrollo. La API y la web deben exponerse mediante URLs HTTPS alcanzables; MariaDB y S3 permanecen privados.")
  add_callout(document, "Limitación", "Esta modalidad depende de que la computadora y los túneles estén activos. Es apropiada para validación técnica, no para exámenes reales o una prueba de carga representativa.", "warning")

  document.add_heading("19.3 Comandos de referencia", level=2)
  add_table(document, ["Objetivo", "Comando"], [
    ("Migrar base", "pnpm --filter @proctoring/api migrate"),
    ("Iniciar API en desarrollo", "pnpm api:dev"),
    ("Iniciar web en desarrollo", "pnpm web:dev"),
    ("Ejecutar pruebas", "pnpm test"),
    ("Compilar web", "pnpm web:build"),
    ("Eliminar evidencia vencida", "pnpm api:evidence:purge"),
    ("Generar fixtures", "pnpm pilot:fixtures"),
    ("Ejecutar carga", "pnpm pilot:load"),
  ], [3300, 6060])

  document.add_heading("19.4 Criterio final de aceptación", level=2)
  add_paragraph(document, "El sistema está preparado para un piloto únicamente cuando Moodle, API, MariaDB, almacenamiento, panel y dispositivos han sido validados en el entorno de destino; no existe acceso entre cursos; cada acceso a evidencia queda auditado; las alertas son revisadas por personas; y ninguna función modifica automáticamente una calificación.")

  document.core_properties.title = "Manual de implementación y uso - Proctoring UDS"
  document.core_properties.subject = "Implementación de MariaDB, API, panel, Moodle y operación de proctoring"
  document.core_properties.author = "Proctoring UDS"
  document.core_properties.keywords = "Moodle, proctoring, MariaDB, implementación, manual de usuario"
  OUTPUT.parent.mkdir(parents=True, exist_ok=True)
  document.save(OUTPUT)
  return OUTPUT


if __name__ == "__main__":
  print(build_document())
